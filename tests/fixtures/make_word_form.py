"""
Generates flat_word_table_form.pdf: a form typed in a word processor (a ruled
table, underscores, a one-cell box and ☐ tick boxes) and exported to PDF.

Development only; needs python-docx and LibreOffice:

    python tests/fixtures/make_word_form.py

This software is released under the AGPL-3.0 license.
"""

import os
import subprocess
import tempfile

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))

if __name__ == "__main__":
    d = Document()
    d.add_heading("Physiotherapy Referral", 1)
    d.add_paragraph("Please complete all sections and fax to the clinic.")
    table = d.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    for i, label in enumerate(["Patient name", "Date of birth", "Health card number", "Phone",
                               "Referring clinician", "Date of referral"]):
        table.cell(i, 0).text = label
    d.add_paragraph("")
    d.add_paragraph("Diagnosis: ________________________________________________")
    d.add_paragraph("Reason for referral / goals of treatment:")
    box = d.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.rows[0].height = Pt(70)
    d.add_paragraph("")
    d.add_paragraph("Urgency:   ☐ Routine     ☐ Urgent (within 1 week)")
    d.add_paragraph("Precautions: ☐ None   ☐ Weight bearing restrictions   ☐ Other")
    d.add_paragraph("")
    d.add_paragraph("Signature: ______________________________     Date: _______________")

    with tempfile.TemporaryDirectory() as work:
        docx_path = os.path.join(work, "flat_word_table_form.docx")
        d.save(docx_path)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, docx_path], check=True)
    print("written")
